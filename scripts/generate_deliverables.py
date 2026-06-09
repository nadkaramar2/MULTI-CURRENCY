#!/usr/bin/env python3
"""Generate PPTX, PDF executive brief, and OpenAPI catalog for MULTI-CURRENCY project."""

import os
import re
import json
from pathlib import Path
from datetime import date

ROOT = Path(r"C:\MULTI-CURRENCY1")
BACKEND = ROOT / "MultiCurrencyCodeBackup" / "SequroCMSAccountManagementService_app"
OUT = ROOT / "deliverables"
OUT.mkdir(exist_ok=True)

CONTEXT_PATH = "/AccountManagementAPI"
BASE_URL = "http://localhost:8285" + CONTEXT_PATH


def scan_controllers():
    """Parse Java controllers for RequestMapping endpoints."""
    java_root = BACKEND / "src" / "main" / "java"
    controllers = []
    class_base_re = re.compile(r'@RequestMapping\s*\(\s*(?:value\s*=\s*)?["\']([^"\']+)["\']')
    method_re = re.compile(
        r'@RequestMapping\s*\(\s*(?:value\s*=\s*|path\s*=\s*)?["\']([^"\']+)["\']'
        r'(?:\s*,\s*method\s*=\s*RequestMethod\.(\w+))?',
        re.MULTILINE,
    )
    post_mapping = re.compile(r'@PostMapping\s*\(\s*(?:value\s*=\s*)?["\']([^"\']+)["\']')
    get_mapping = re.compile(r'@GetMapping\s*\(\s*(?:value\s*=\s*)?["\']([^"\']+)["\']')
    class_name_re = re.compile(r'public\s+class\s+(\w+)')

    for path in sorted(java_root.rglob("*Controller*.java")):
        text = path.read_text(encoding="utf-8", errors="ignore")
        if "@RestController" not in text and "@Controller" not in text:
            continue
        cls = class_name_re.search(text)
        class_name = cls.group(1) if cls else path.stem
        base = ""
        m = class_base_re.search(text)
        if m:
            base = m.group(1).strip()
            if not base.startswith("/"):
                base = "/" + base

        rel = str(path.relative_to(java_root)).replace("\\", "/")
        pkg = "api" if "/api/controller/" in rel.replace("\\", "/") else "admin"

        endpoints = []
        for mm in method_re.finditer(text):
            sub = mm.group(1)
            method = (mm.group(2) or "POST").upper()
            full = join_path(base, sub)
            endpoints.append({"method": method, "path": full, "operation_id": op_id(class_name, sub, method)})

        for mm in post_mapping.finditer(text):
            sub = mm.group(1)
            full = join_path(base, sub)
            endpoints.append({"method": "POST", "path": full, "operation_id": op_id(class_name, sub, "POST")})

        for mm in get_mapping.finditer(text):
            sub = mm.group(1)
            full = join_path(base, sub)
            endpoints.append({"method": "GET", "path": full, "operation_id": op_id(class_name, sub, "GET")})

        # Dedupe
        seen = set()
        unique = []
        for e in endpoints:
            k = (e["method"], e["path"])
            if k not in seen:
                seen.add(k)
                unique.append(e)

        if unique or base:
            controllers.append({
                "controller": class_name,
                "file": rel,
                "package": pkg,
                "base_path": base or "/",
                "endpoints": sorted(unique, key=lambda x: x["path"]),
            })
    return controllers


def join_path(base, sub):
    if not sub:
        return base or "/"
    if sub.startswith("/") and not base:
        return sub
    if sub.startswith("/"):
        return (base.rstrip("/") + sub) if base else sub
    return (base.rstrip("/") + "/" + sub) if base else "/" + sub


def op_id(class_name, sub, method):
    name = re.sub(r"[^a-zA-Z0-9]", "_", (class_name + "_" + sub).strip("_"))
    return f"{method.lower()}_{name}"[:80]


def write_openapi(controllers):
    paths = {}
    tags = {}
    for ctrl in controllers:
        tag = ctrl["controller"]
        tags[tag] = {"name": tag, "description": f"From {ctrl['file']}"}
        for ep in ctrl["endpoints"]:
            p = ep["path"]
            if p not in paths:
                paths[p] = {}
            paths[p][ep["method"].lower()] = {
                "tags": [tag],
                "operationId": ep["operation_id"],
                "summary": f"{ep['method']} {p}",
                "description": f"Controller: {ctrl['controller']}",
                "responses": {
                    "200": {
                        "description": "ProcessResponse envelope (code, message, data)",
                        "content": {
                            "application/json": {
                                "schema": {"$ref": "#/components/schemas/ProcessResponse"}
                            }
                        },
                    }
                },
                "requestBody": {
                    "required": True,
                    "content": {
                        "application/json": {
                            "schema": {"type": "object", "description": "JSON body or PayloadReqRes encrypted envelope"}
                        }
                    },
                } if ep["method"] in ("POST", "PUT", "PATCH") else None,
            }
            if paths[p][ep["method"].lower()].get("requestBody") is None:
                del paths[p][ep["method"].lower()]["requestBody"]

    spec = {
        "openapi": "3.0.3",
        "info": {
            "title": "MULTI-CURRENCY Account Management API",
            "description": "Auto-generated from Spring @RestController scan. Context path: " + CONTEXT_PATH,
            "version": "1.0.0",
            "contact": {"name": "Sequro CMS AMS"},
        },
        "servers": [{"url": BASE_URL, "description": "Local Spring Boot"}],
        "tags": list(tags.values()),
        "paths": dict(sorted(paths.items())),
        "components": {
            "schemas": {
                "ProcessResponse": {
                    "type": "object",
                    "properties": {
                        "code": {"type": "string", "example": "S0000"},
                        "message": {"type": "string"},
                        "data": {"type": "object"},
                        "status": {"type": "string"},
                    },
                },
                "PayloadReqRes": {
                    "type": "object",
                    "description": "Encrypted payload wrapper for mobile APIs",
                    "properties": {
                        "payload": {"type": "string"},
                    },
                },
            },
            "securitySchemes": {
                "bearerAuth": {"type": "http", "scheme": "bearer", "bearerFormat": "JWT"},
                "apiKey": {"type": "apiKey", "in": "header", "name": "apikey"},
            },
        },
        "security": [{"bearerAuth": []}, {"apiKey": []}],
    }

    yaml_path = OUT / "MULTI-CURRENCY-openapi.yaml"
    json_path = OUT / "MULTI-CURRENCY-openapi.json"
    md_path = OUT / "MULTI-CURRENCY-API-CATALOG.md"

    try:
        import yaml
        yaml_path.write_text(yaml.dump(spec, sort_keys=False, allow_unicode=True, default_flow_style=False), encoding="utf-8")
    except ImportError:
        yaml_path.write_text(json.dumps(spec, indent=2), encoding="utf-8")

    json_path.write_text(json.dumps(spec, indent=2), encoding="utf-8")

    # Markdown catalog
    lines = [
        "# MULTI-CURRENCY API Catalog",
        "",
        f"**Generated:** {date.today().isoformat()}  ",
        f"**Base URL:** `{BASE_URL}`  ",
        f"**Controllers scanned:** {len(controllers)}  ",
        f"**Total endpoints:** {sum(len(c['endpoints']) for c in controllers)}  ",
        "",
        "---",
        "",
    ]
    for ctrl in controllers:
        lines.append(f"## {ctrl['controller']} (`{ctrl['package']}`)")
        lines.append(f"**File:** `{ctrl['file']}`  ")
        lines.append(f"**Base path:** `{ctrl['base_path']}`  ")
        lines.append("")
        lines.append("| Method | Path | Operation ID |")
        lines.append("|--------|------|--------------|")
        for ep in ctrl["endpoints"]:
            lines.append(f"| {ep['method']} | `{ep['path']}` | `{ep['operation_id']}` |")
        lines.append("")
    md_path.write_text("\n".join(lines), encoding="utf-8")
    return yaml_path, json_path, md_path


def generate_pptx():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    NAVY = RGBColor(0x0B, 0x1F, 0x3A)
    GOLD = RGBColor(0xC9, 0xA2, 0x27)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    GRAY = RGBColor(0x4A, 0x55, 0x68)

    slides_data = [
        ("title", "MULTI-CURRENCY", "Enterprise Account & Wallet Management Platform",
         ["Sequro CMS · Account Management System", "One platform. Multiple currencies. Complete control."]),
        ("content", "Product Overview", None, [
            "End-to-end Account Management System for banks, fintechs & NBFCs",
            "Native multi-currency wallet per customer account",
            "Unified admin console + mobile-ready encrypted APIs",
            "Solves: fragmented FX ops, slow onboarding, weak audit trails",
        ]),
        ("content", "Key Features", None, [
            "Multi-currency wallets (USD, EUR, INR, NGN…)",
            "FX conversion engine with configurable rates",
            "Load money & cross-currency transfers",
            "GL integration, bulk pay, credit card lifecycle",
            "Role-based admin UI · Maker-checker workflows",
        ]),
        ("content", "Multi-Currency Workflow", None, [
            "① Configure currencies & FX rates",
            "② Open account → auto-create currency wallets",
            "③ Wallet ID = Account Number + Currency Code",
            "④ Load → Transfer → Statements & regulatory reports",
            "Each currency is a first-class wallet with balance & priority",
        ]),
        ("content", "Security Features", None, [
            "JWT authentication · API secret key validation",
            "AES encrypted payloads for mobile channel",
            "Customer PIN on sensitive wallet operations",
            "Maker-checker: dormancy, bulk, journal transfers",
            "Full txn request/response audit · Participant isolation",
        ]),
        ("two_col", "Technology Stack", None, [
            ("Backend", "Java 8 · Spring Boot 2.7 · Hibernate · MySQL"),
            ("Frontend", "React 18 · Redux · Axios · Tailwind"),
            ("Security", "Spring Security · JWT · AES encryption"),
            ("Integration", "REST · Middleware · SMS/Email · QR"),
        ]),
        ("two_col", "Business Benefits", None, [
            ("CEO / Business", "Faster MC product launch · FX fee revenue"),
            ("Operations", "100+ ops in one console"),
            ("Compliance", "LRS/TCS · dormancy · closure · txn logs"),
            ("IT / Customer", "Monolith clarity · seamless mobile wallets"),
        ]),
        ("content", "Why Choose Our Solution", None, [
            "✓ 80+ REST controllers · 95+ entities · 120+ admin screens",
            "✓ Multi-currency by design (wallet master, conversion, transfer)",
            "✓ Full banking: accounts, GL, bulk, cards, tiers, complaints",
            "✓ Dual channel: operations UI + encrypted mobile APIs",
            "✓ Multi-tenant participant model for BaaS & group banks",
        ]),
        ("content", "Future Enhancements", None, [
            "Real-time FX feed · Open Banking / ISO 20022",
            "API-layer RBAC hardening · Cloud-native (K8s)",
            "Treasury analytics dashboard · Digital eKYC onboarding",
            "Incremental modernization — no rip-and-replace",
        ]),
        ("title", "Thank You", "Schedule Your Live Demo", [
            "onboarding → wallet creation → FX transfer → statement",
            "MULTI-CURRENCY · The OS for multi-currency account management",
        ]),
    ]

    def set_title_style(shape, size=36, color=NAVY, bold=True):
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(size)
                r.font.bold = bold
                r.font.color.rgb = color

    for stype, title, subtitle, bullets in slides_data:
        if stype == "title":
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
            bg.fill.solid()
            bg.fill.fore_color.rgb = NAVY
            bg.line.fill.background()
            tb = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.5), Inches(1.2))
            tf = tb.text_frame
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(44)
            p.font.bold = True
            p.font.color.rgb = WHITE
            if subtitle:
                tb2 = slide.shapes.add_textbox(Inches(0.8), Inches(3.5), Inches(11.5), Inches(1))
                p2 = tb2.text_frame.paragraphs[0]
                p2.text = subtitle
                p2.font.size = Pt(22)
                p2.font.color.rgb = GOLD
            if bullets:
                tb3 = slide.shapes.add_textbox(Inches(0.8), Inches(4.8), Inches(11), Inches(2))
                for i, b in enumerate(bullets):
                    para = tb3.text_frame.paragraphs[0] if i == 0 else tb3.text_frame.add_paragraph()
                    para.text = b
                    para.font.size = Pt(16)
                    para.font.color.rgb = WHITE
        else:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.1))
            bar.fill.solid()
            bar.fill.fore_color.rgb = NAVY
            bar.line.fill.background()
            tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.25), Inches(12), Inches(0.7))
            tb.text_frame.paragraphs[0].text = title
            tb.text_frame.paragraphs[0].font.size = Pt(28)
            tb.text_frame.paragraphs[0].font.bold = True
            tb.text_frame.paragraphs[0].font.color.rgb = WHITE

            if stype == "two_col" and bullets:
                left, right = 0, len(bullets) // 2 + len(bullets) % 2
                for col, (x, items) in enumerate([(0.6, bullets[:2]), (6.8, bullets[2:])]):
                    box = slide.shapes.add_textbox(Inches(x), Inches(1.5), Inches(5.5), Inches(5))
                    tf = box.text_frame
                    for i, (head, body) in enumerate(items):
                        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                        p.text = head
                        p.font.bold = True
                        p.font.size = Pt(18)
                        p.font.color.rgb = NAVY
                        p2 = tf.add_paragraph()
                        p2.text = body
                        p2.font.size = Pt(14)
                        p2.font.color.rgb = GRAY
                        p2.space_after = Pt(12)
            elif bullets:
                box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.5), Inches(5.5))
                tf = box.text_frame
                for i, b in enumerate(bullets):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.text = ("• " + b) if not b.startswith("✓") else b
                    p.font.size = Pt(20)
                    p.font.color.rgb = GRAY
                    p.space_after = Pt(10)

    out = OUT / "MULTI-CURRENCY-Client-Presentation.pptx"
    prs.save(str(out))
    return out


def generate_pdf():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    out = OUT / "MULTI-CURRENCY-Executive-Brief.pdf"
    doc = SimpleDocTemplate(str(out), pagesize=A4, leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    navy = colors.HexColor("#0B1F3A")
    gold = colors.HexColor("#C9A227")

    title_style = ParagraphStyle("Title", parent=styles["Heading1"], fontSize=22, textColor=navy, spaceAfter=6)
    sub_style = ParagraphStyle("Sub", parent=styles["Normal"], fontSize=11, textColor=gold, spaceAfter=14)
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=12, textColor=navy, spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle("Body", parent=styles["Normal"], fontSize=9, leading=12, spaceAfter=4)

    story = [
        Paragraph("MULTI-CURRENCY — Executive Brief", title_style),
        Paragraph("Sequro CMS Account Management System · Multi-Currency Edition", sub_style),
        Paragraph(f"<b>Date:</b> {date.today().strftime('%B %d, %Y')}", body),
        Spacer(1, 0.3 * cm),
        Paragraph("<b>Elevator Pitch</b>", h2),
        Paragraph(
            "MULTI-CURRENCY gives financial institutions one platform to onboard customers, "
            "operate per-currency wallets with live FX rates, move money across currencies and channels, "
            "and stay audit-ready—with GL posting, maker-checker workflows, and a modern operations console.",
            body,
        ),
        Paragraph("<b>At a Glance</b>", h2),
    ]

    data = [
        ["Component", "Technology", "Scale"],
        ["Backend API", "Spring Boot 2.7 · Java 8 · Hibernate · MySQL", "~80 controllers · ~95 entities"],
        ["Admin UI", "React 18 · Redux · Axios", "~120 screens · RBAC menus"],
        ["Security", "JWT · API key · AES encryption", "Mobile + admin channels"],
        ["Core domain", "Multi-currency wallets", "Wallet ID = Account + Currency"],
    ]
    t = Table(data, colWidths=[3.2 * cm, 7 * cm, 5 * cm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), navy),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6F8")]),
    ]))
    story.extend([t, Spacer(1, 0.2 * cm)])

    for heading, items in [
        ("Business Benefits", [
            "Launch multi-currency products faster; monetize FX and wallet fees",
            "Single operations console for configuration, transfers, and reports",
            "LRS/TCS, dormancy, closure, and full transaction audit trails",
            "Participant-based multi-tenancy for BaaS and banking groups",
        ]),
        ("Key Modules (Implemented)", [
            "MultiCurrencyWalletAccountController — wallets, priority, FX lookup",
            "LoadMoneyController — load money & transferCurrency",
            "AccountMasterController — customer account lifecycle",
            "BulkTransferController · JournalTransfer · SignUp/SignIn mobile APIs",
        ]),
        ("Recommended Demo Path", [
            "Login → Currency master → FX conversion → Create account → Load money report",
        ]),
    ]:
        story.append(Paragraph(f"<b>{heading}</b>", h2))
        for item in items:
            story.append(Paragraph(f"• {item}", body))

    story.append(Spacer(1, 0.3 * cm))
    story.append(Paragraph(
        "<i>Contact: Schedule a live demo — onboarding → wallet creation → FX transfer → statement</i>",
        ParagraphStyle("Footer", parent=body, fontSize=8, textColor=colors.grey),
    ))

    doc.build(story)
    return out


def main():
    print("Scanning controllers...")
    controllers = scan_controllers()
    total_eps = sum(len(c["endpoints"]) for c in controllers)
    print(f"Found {len(controllers)} controllers, {total_eps} endpoints")

    print("Generating OpenAPI catalog...")
    y, j, m = write_openapi(controllers)
    print(f"  {y}")
    print(f"  {j}")
    print(f"  {m}")

    print("Generating PowerPoint...")
    pptx = generate_pptx()
    print(f"  {pptx}")

    print("Generating PDF executive brief...")
    pdf = generate_pdf()
    print(f"  {pdf}")

    print("Done.")


if __name__ == "__main__":
    main()
